import pdfplumber
from docx import Document
import re
import os
from typing import Tuple, Dict, List, Optional
from PIL import Image
import pytesseract
import cv2
import numpy as np

class FileParser:
    """Handle PDF, Word, Text, and Image file parsing"""
    
    # Supported image formats
    IMAGE_FORMATS = ['jpg', 'jpeg', 'png', 'gif', 'bmp', 'tiff', 'webp']
    
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> Tuple[str, int]:
        """Extract text from PDF file and return text and page count"""
        text = ""
        page_count = 0
        try:
            with pdfplumber.open(file_path) as pdf:
                page_count = len(pdf.pages)
                for page_num, page in enumerate(pdf.pages, 1):
                    # Extract regular text
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n[PAGE {page_num}]\n{page_text}\n"
                    
                    # Extract tables safely
                    try:
                        tables = page.extract_tables()
                        if tables:
                            for table in tables:
                                for row in table:
                                    row_text = " | ".join(str(cell) for cell in row if cell)
                                    if row_text:
                                        text += row_text + "\n"
                    except:
                        pass
        except Exception as e:
            raise Exception(f"PDF extraction error: {str(e)}")
        return text, page_count
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from Word document"""
        text = ""
        try:
            doc = Document(file_path)
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text += paragraph.text + "\n"
            
            # Extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text:
                            row_text.append(cell.text.strip())
                    if row_text:
                        text += " | ".join(row_text) + "\n"
        except Exception as e:
            raise Exception(f"DOCX extraction error: {str(e)}")
        return text
    
    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        """Extract text from text file"""
        encodings = ['utf-8', 'latin-1', 'cp1252', 'iso-8859-1']
        
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as file:
                    return file.read()
            except UnicodeDecodeError:
                continue
        
        raise Exception("Could not read text file with any encoding")
    
    @staticmethod
    def extract_text_from_image(file_path: str, preprocess: bool = True) -> str:
        """Extract text from image using OCR"""
        try:
            # Open image
            img = Image.open(file_path)
            
            if preprocess:
                # Convert PIL to OpenCV for preprocessing
                img_cv = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
                
                # Preprocess for better OCR
                gray = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)
                
                # Apply thresholding
                _, thresh = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
                
                # Denoise
                denoised = cv2.medianBlur(thresh, 3)
                
                # Convert back to PIL
                img = Image.fromarray(denoised)
            
            # Extract text using Tesseract
            # Use multiple languages if needed
            text = pytesseract.image_to_string(img, lang='eng')
            
            # Clean the extracted text
            text = re.sub(r'\s+', ' ', text)
            text = re.sub(r'[^\w\s\.\?\!,;:\-\(\)]', '', text)
            
            return text.strip()
            
        except Exception as e:
            raise Exception(f"Image OCR extraction error: {str(e)}")
    
    @staticmethod
    def extract_text_from_image_with_details(file_path: str) -> Dict:
        """Extract text with confidence scores and layout"""
        try:
            img = Image.open(file_path)
            
            # Get detailed OCR data
            data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
            
            text_blocks = []
            full_text = ""
            
            for i, word in enumerate(data['text']):
                if word.strip() and int(data['conf'][i]) > 30:
                    text_blocks.append({
                        'word': word,
                        'confidence': float(data['conf'][i]),
                        'bbox': {
                            'x': data['left'][i],
                            'y': data['top'][i],
                            'width': data['width'][i],
                            'height': data['height'][i]
                        },
                        'line_num': data['line_num'][i],
                        'block_num': data['block_num'][i]
                    })
                    full_text += word + " "
            
            # Calculate average confidence
            confidences = [int(conf) for conf in data['conf'] if conf != '-1']
            avg_confidence = np.mean(confidences) if confidences else 0
            
            return {
                'text': full_text.strip(),
                'words': text_blocks,
                'avg_confidence': avg_confidence,
                'word_count': len(text_blocks),
                'line_count': len(set(b['line_num'] for b in text_blocks)),
                'has_numbers': any(re.search(r'\d', w['word']) for w in text_blocks)
            }
            
        except Exception as e:
            raise Exception(f"Detailed image OCR failed: {str(e)}")
    
    @staticmethod
    def extract_text_from_scanned_pdf(file_path: str) -> Tuple[str, int]:
        """Extract text from scanned PDF using OCR"""
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise Exception("PyMuPDF (fitz) is required for scanned PDF OCR. Install: pip install PyMuPDF")
        
        text = ""
        page_count = 0
        
        try:
            doc = fitz.open(file_path)
            page_count = len(doc)
            
            for page_num, page in enumerate(doc, 1):
                # Get image from page
                pix = page.get_pixmap()
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                
                # Extract text using OCR
                page_text = pytesseract.image_to_string(img)
                if page_text:
                    text += f"\n[PAGE {page_num}]\n{page_text}\n"
            
            doc.close()
            
        except Exception as e:
            raise Exception(f"Scanned PDF OCR error: {str(e)}")
        
        return text, page_count
    
    @staticmethod
    def is_image_file(file_path: str) -> bool:
        """Check if file is an image based on extension"""
        ext = os.path.splitext(file_path)[1].lower().replace('.', '')
        return ext in FileParser.IMAGE_FORMATS
    
    @staticmethod
    def is_scanned_pdf(file_path: str) -> bool:
        """Check if PDF is scanned (no selectable text)"""
        try:
            import fitz
            doc = fitz.open(file_path)
            # Check if first page has text
            page = doc[0]
            text = page.get_text()
            doc.close()
            return len(text.strip()) == 0
        except:
            return False
    
    @staticmethod
    def parse_file(file_path: str, file_type: str) -> Tuple[str, Dict]:
        """Main method to parse files and extract metadata"""
        page_count = 0
        is_scanned = False
        
        # Handle image files
        if file_type in ['JPG', 'JPEG', 'PNG', 'GIF', 'BMP', 'TIFF', 'WEBP']:
            text = FileParser.extract_text_from_image(file_path)
            metadata = {
                'word_count': len(text.split()),
                'char_count': len(text),
                'line_count': len(text.split('\n')),
                'file_size': os.path.getsize(file_path),
                'page_count': 1,
                'file_type': 'image',
                'ocr_used': True
            }
            return text, metadata
        
        # Handle PDF
        elif file_type == 'PDF':
            # Check if PDF is scanned (no selectable text)
            try:
                import fitz
                doc = fitz.open(file_path)
                first_page_text = doc[0].get_text()
                doc.close()
                is_scanned = len(first_page_text.strip()) == 0
            except:
                is_scanned = False
            
            if is_scanned:
                # Use OCR for scanned PDF
                text, page_count = FileParser.extract_text_from_scanned_pdf(file_path)
                ocr_used = True
            else:
                # Regular PDF extraction
                text, page_count = FileParser.extract_text_from_pdf(file_path)
                ocr_used = False
            
            metadata = {
                'word_count': len(text.split()),
                'char_count': len(text),
                'line_count': len(text.split('\n')),
                'file_size': os.path.getsize(file_path),
                'page_count': page_count,
                'ocr_used': ocr_used
            }
            return text, metadata
        
        # Handle DOCX
        elif file_type == 'DOCX':
            text = FileParser.extract_text_from_docx(file_path)
            metadata = {
                'word_count': len(text.split()),
                'char_count': len(text),
                'line_count': len(text.split('\n')),
                'file_size': os.path.getsize(file_path),
                'page_count': 0,
                'ocr_used': False
            }
            return text, metadata
        
        # Handle TXT
        elif file_type == 'TXT':
            text = FileParser.extract_text_from_txt(file_path)
            metadata = {
                'word_count': len(text.split()),
                'char_count': len(text),
                'line_count': len(text.split('\n')),
                'file_size': os.path.getsize(file_path),
                'page_count': 0,
                'ocr_used': False
            }
            return text, metadata
        
        else:
            raise Exception(f"Unsupported file type: {file_type}")
    
    @staticmethod
    def detect_sections(text: str) -> Dict[str, str]:
        """Detect common sections in educational content"""
        sections = {}
        
        patterns = {
            'questions': r'(?:Questions?|Exercises?|Problems?)(?:\s*:)?\s*\n(.*?)(?=\n\s*(?:[A-Z][a-z]+:|$))',
            'examples': r'(?:Examples?|Illustrations?)(?:\s*:)?\s*\n(.*?)(?=\n\s*(?:[A-Z][a-z]+:|$))',
            'answers': r'(?:Answers?|Solutions?)(?:\s*:)?\s*\n(.*?)(?=\n\s*(?:[A-Z][a-z]+:|$))',
            'summary': r'(?:Summary|Conclusion|Key Points)(?:\s*:)?\s*\n(.*?)(?=\n\s*(?:[A-Z][a-z]+:|$))',
            'objectives': r'(?:Objectives|Learning Outcomes|Goals)(?:\s*:)?\s*\n(.*?)(?=\n\s*(?:[A-Z][a-z]+:|$))',
        }
        
        for section, pattern in patterns.items():
            match = re.search(pattern, text, re.DOTALL | re.IGNORECASE)
            if match:
                sections[section] = match.group(1).strip()
        
        return sections